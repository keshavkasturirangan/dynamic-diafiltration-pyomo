from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/kkasturi/GitHub/keshav-dev-dynamic-diafiltration-pyomo")
RUN_ROOT = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260408-codex-paper-check"
OUTDIR = RUN_ROOT / "ppt_full_comparison"
FIGDIR = RUN_ROOT / "figures"
EXACTDIR = RUN_ROOT / "figures_exact"
PUB_D1 = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260306-023356-paper-pdf-extract" / "pdf_extract" / "data1_main" / "images"
PUB_D2 = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260306-023356-paper-pdf-extract" / "pdf_extract" / "data2_main" / "images"


def add_title(slide, text, *, top=0.2, size=24):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.bold = True
    run.font.size = Pt(size)


def add_text(slide, text, *, left=0.55, top=6.85, width=12.0, height=0.35, size=13):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)


def add_label(slide, text, *, left, top):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(2.8), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(14)
    run.font.bold = True


def vstack(paths, out_path, *, bg="white", gap=30):
    images = [Image.open(path).convert("RGB") for path in paths]
    width = max(img.width for img in images)
    height = sum(img.height for img in images) + gap * (len(images) - 1)
    canvas = Image.new("RGB", (width, height), bg)
    y = 0
    for img in images:
        x = (width - img.width) // 2
        canvas.paste(img, (x, y))
        y += img.height + gap
    canvas.save(out_path)
    return out_path


def add_side_by_side(slide, published, regenerated):
    add_label(slide, "Published", left=0.6, top=0.78)
    add_label(slide, "Unified Regenerated", left=6.95, top=0.78)
    slide.shapes.add_picture(str(published), Inches(0.45), Inches(1.05), width=Inches(5.8))
    slide.shapes.add_picture(str(regenerated), Inches(6.85), Inches(1.05), width=Inches(5.95))


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)

    d2_img008 = vstack(
        [FIGDIR / "Js_predict0.png", FIGDIR / "Jw_predict.png", FIGDIR / "Js_predict1.png"],
        OUTDIR / "data2_img008_unified_stack.png",
    )
    d2_img010 = vstack(
        [FIGDIR / "startup_barplot.png", FIGDIR / "concentrating_residuals_boxplot.png"],
        OUTDIR / "data2_img010_unified_stack.png",
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = [
        {
            "title": "DATA2 Figure 7",
            "published": PUB_D2 / "img-007.png",
            "regenerated": EXACTDIR / "data2_img007_unified_exact.png",
            "note": "Published figure on the left and current unified-codebase regeneration on the right.",
        },
        {
            "title": "DATA2 Figure 8",
            "published": PUB_D2 / "img-008.png",
            "regenerated": d2_img008,
            "note": "Unified workflow reproduces the three published panels using the current DATA2 artifact path.",
        },
        {
            "title": "DATA2 Figure 9",
            "published": PUB_D2 / "img-009.png",
            "regenerated": FIGDIR / "partition_sensitivity.png",
            "note": "Unified workflow reproduces the partition-sensitivity contour figure.",
        },
        {
            "title": "DATA2 Figure 10",
            "published": PUB_D2 / "img-010.png",
            "regenerated": d2_img010,
            "note": "Unified workflow reproduces the startup-improvement panel and residuals boxplot.",
        },
        {
            "title": "DATA1 Figure 2",
            "published": PUB_D1 / "img-004.jpg",
            "regenerated": EXACTDIR / "data1_img004_unified_exact.png",
            "note": "Unified workflow reproduces the four-panel mass and concentration comparison figure.",
        },
        {
            "title": "DATA1 Figure 3",
            "published": PUB_D1 / "img-005.jpg",
            "regenerated": FIGDIR / "concentration_range.png",
            "note": "Unified workflow reproduces the experiment-space concentration-range figure.",
        },
        {
            "title": "DATA1 Figure 4",
            "published": PUB_D1 / "img-006.jpg",
            "regenerated": EXACTDIR / "data1_img006_unified_exact.png",
            "note": "Unified workflow reproduces the sigma-sensitivity figure across mass, permeate, and retentate panels.",
        },
        {
            "title": "DATA1 Figure 5",
            "published": PUB_D1 / "img-007.jpg",
            "regenerated": EXACTDIR / "data1_img007_unified_exact.png",
            "note": "Unified workflow reproduces the fixed-sigma contour maps over sigma and Lp-style sweeps.",
        },
        {
            "title": "DATA1 Figure 6",
            "published": PUB_D1 / "img-008.jpg",
            "regenerated": EXACTDIR / "data1_img008_unified_exact.png",
            "note": "Unified workflow reproduces the fixed-B contour maps over B and Lp sweeps.",
        },
    ]

    for item in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, item["title"])
        add_side_by_side(slide, item["published"], item["regenerated"])
        add_text(slide, item["note"])

    prs.save(str(OUTDIR / "published_vs_unified_full_comparison.pptx"))


if __name__ == "__main__":
    main()
