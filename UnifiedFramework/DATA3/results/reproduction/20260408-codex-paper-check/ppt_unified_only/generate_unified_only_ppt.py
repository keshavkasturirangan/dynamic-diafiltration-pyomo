from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/kkasturi/GitHub/keshav-dev-dynamic-diafiltration-pyomo")
OUTDIR = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260408-codex-paper-check" / "ppt_unified_only"
FIGDIR = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260408-codex-paper-check" / "figures"
PUB_D2 = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260306-023356-paper-pdf-extract" / "pdf_extract" / "data2_main" / "images"
PUB_D1 = ROOT / "UnifiedFramework" / "DATA3" / "results" / "reproduction" / "20260306-023356-paper-pdf-extract" / "pdf_extract" / "data1_main" / "images"


def add_title(slide, text, left=0.4, top=0.2, width=12.5, height=0.5, size=24):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    p.alignment = PP_ALIGN.LEFT


def add_body(slide, text, left, top, width, height, size=14):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)


def add_label(slide, text, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(14)
    run.font.bold = True


def add_one_vs_one(prs, title, published, regenerated, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    add_label(slide, "Published", 0.6, 0.75, 2.0, 0.3)
    add_label(slide, "Unified Regenerated", 6.8, 0.75, 3.2, 0.3)
    slide.shapes.add_picture(str(published), Inches(0.5), Inches(1.1), width=Inches(5.8))
    slide.shapes.add_picture(str(regenerated), Inches(6.7), Inches(1.1), width=Inches(6.0))
    add_body(slide, caption, 0.6, 6.7, 12.0, 0.5, size=13)


def add_one_vs_many(prs, title, published, regenerated, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    add_label(slide, "Published", 0.6, 0.75, 2.0, 0.3)
    add_label(slide, "Unified Regenerated", 7.1, 0.75, 3.2, 0.3)
    slide.shapes.add_picture(str(published), Inches(0.4), Inches(1.0), width=Inches(6.0))
    top = 1.0
    for path in regenerated:
        slide.shapes.add_picture(str(path), Inches(7.0), Inches(top), width=Inches(5.8))
        top += 1.8
    add_body(slide, caption, 0.6, 6.9, 12.0, 0.4, size=13)


def add_missing_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Published Figures Still Missing Exact Unified Matches")
    lines = [
        "DATA2 img-007.png",
        "DATA1 img-004.jpg",
        "DATA1 img-006.jpg",
        "DATA1 img-007.jpg",
        "DATA1 img-008.jpg",
        "",
        "These should not be presented as unified-regenerated yet.",
        "Only the four comparison slides in this deck are exact unified-output matches.",
    ]
    add_body(slide, "\n".join(lines), 0.8, 1.2, 6.2, 4.6, size=20)


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Unified-Only Published Figure Comparison", top=0.5, size=28)
    add_body(
        slide,
        "This deck includes only figures whose regenerated counterparts come from the current unified workflow.\n\n"
        "Included:\n"
        "- DATA2 img-008, img-009, img-010\n"
        "- DATA1 img-005\n\n"
        "Excluded for now:\n"
        "- DATA2 img-007\n"
        "- DATA1 img-004, img-006, img-007, img-008",
        0.8,
        1.5,
        8.5,
        4.8,
        size=22,
    )

    add_one_vs_many(
        prs,
        "DATA2 Published img-008 vs Unified Regenerated",
        PUB_D2 / "img-008.png",
        [FIGDIR / "Js_predict0.png", FIGDIR / "Jw_predict.png", FIGDIR / "Js_predict1.png"],
        "Unified output reproduces the three published panels: concentration vs time, Jw vs time, and Js vs cin,f.",
    )
    add_one_vs_one(
        prs,
        "DATA2 Published img-009 vs Unified Regenerated",
        PUB_D2 / "img-009.png",
        FIGDIR / "partition_sensitivity.png",
        "Unified output reproduces the partition-sensitivity contour figure.",
    )
    add_one_vs_many(
        prs,
        "DATA2 Published img-010 vs Unified Regenerated",
        PUB_D2 / "img-010.png",
        [FIGDIR / "startup_barplot.png", FIGDIR / "concentrating_residuals_boxplot.png"],
        "Unified output reproduces the startup-information panel and the concentrating residuals boxplot.",
    )
    add_one_vs_one(
        prs,
        "DATA1 Published img-005 vs Unified Regenerated",
        PUB_D1 / "img-005.jpg",
        FIGDIR / "concentration_range.png",
        "Unified output reproduces the concentration-range figure.",
    )
    add_missing_slide(prs)

    prs.save(str(OUTDIR / "unified_only_published_comparison.pptx"))


if __name__ == "__main__":
    main()
