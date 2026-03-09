#!/usr/bin/env python3
"""Build a 10-slide PowerPoint deck from nightly ops documentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for idx, line in enumerate(bullets):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(22)


def add_code_slide(prs: Presentation, title: str, lines: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    x, y, w, h = Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.3)
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()

    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Courier New"
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(40, 40, 40)


def build_deck(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Nightly Ops Validation Runbook"
    prs.core_properties.subject = "Dynamic Diafiltration Validation Automation"

    add_title_slide(
        prs,
        "Nightly Ops Validation",
        "Dynamic Diafiltration: gate, comparison report, and storage-aware prune",
    )

    add_bullets_slide(
        prs,
        "What Runs Nightly",
        [
            "Workflow: .github/workflows/nightly-validation.yml",
            "Status gate over target_validation_status_consolidated.csv",
            "Known non-pass targets handled as warnings from nightly/config/known_nonpass.csv",
            "Digitized paper-vs-unified comparison report is regenerated",
            "Storage prune executes in dry-run archive mode",
        ],
    )

    add_bullets_slide(
        prs,
        "Gate Semantics",
        [
            "Hard failure only on unexpected FAIL or MISSING_VALUE targets",
            "Status mismatches against exception list are treated as failures",
            "NOT_APPLICABLE is never considered a nightly failure",
            "PASS_WITH_EXPLANATION is treated as pass",
            "Goal: stable nightly health signal with intentional warning visibility",
        ],
    )

    add_bullets_slide(
        prs,
        "Nightly Step Order",
        [
            "1) Execute nightly_validation_gate.py",
            "2) Execute compare_digitized_baselines.py (report mode)",
            "3) Execute prune_reproduction_artifacts.py dry-run",
            "4) Review warnings and next-action notes",
            "5) Keep strict failure scope limited to unexpected regression states",
        ],
    )

    add_code_slide(
        prs,
        "Command: Nightly Validation Gate",
        [
            "python UnifiedFramework/DATA3/scripts/validation/nightly/nightly_validation_gate.py \\",
            "  --status-csv UnifiedFramework/DATA3/docs/validation/target_validation_status_consolidated.csv \\",
            "  --exceptions-csv UnifiedFramework/DATA3/docs/validation/nightly/config/known_nonpass.csv",
        ],
    )

    add_code_slide(
        prs,
        "Command: Digitized Comparison Report",
        [
            "python UnifiedFramework/DATA3/scripts/validation/nightly/compare_digitized_baselines.py \\",
            "  --repo-root . \\",
            "  --manifest UnifiedFramework/DATA3/docs/validation/nightly/digitized_baselines/manifest.csv \\",
            "  --thresholds UnifiedFramework/DATA3/docs/validation/nightly/digitized_baselines/thresholds.csv \\",
            "  --out-csv UnifiedFramework/DATA3/docs/validation/nightly/digitized_baselines/comparison_latest.csv",
        ],
    )

    add_code_slide(
        prs,
        "Command: Append Nightly Note",
        [
            "python UnifiedFramework/DATA3/scripts/validation/nightly/append_nightly_note.py \\",
            "  --run-id <run-id> --context <context> --gate-status <PASS|FAIL> \\",
            "  --gate-summary \"No unexpected FAIL/MISSING_VALUE\" \\",
            "  --pytest-status <status> --action <next-step>",
            "",
            "Log target: UnifiedFramework/DATA3/docs/validation/nightly/logs/nightly_test_notes.md",
        ],
    )

    add_code_slide(
        prs,
        "Command: Storage Prune",
        [
            "Dry-run:",
            "python UnifiedFramework/DATA3/scripts/reproduce/prune_reproduction_artifacts.py \\",
            "  --repo-root . --keep-full 3 --mode archive",
            "",
            "Apply:",
            "python UnifiedFramework/DATA3/scripts/reproduce/prune_reproduction_artifacts.py \\",
            "  --repo-root . --keep-full 3 --mode archive --apply",
        ],
    )

    add_bullets_slide(
        prs,
        "Maintaining Known Non-Pass Targets",
        [
            "Edit UnifiedFramework/DATA3/docs/validation/nightly/config/known_nonpass.csv",
            "One row per accepted non-pass target",
            "Set expected_status to FAIL or MISSING_VALUE",
            "Set active=true and maintain review_after date",
            "Remove/deactivate row once target improves to PASS/NOT_APPLICABLE",
        ],
    )

    add_bullets_slide(
        prs,
        "Strict Mode Activation",
        [
            "Enable strict digitized enforcement only when baselines are mature",
            "Optional flags: --fail-on-fail and --fail-on-missing",
            "Activate panel rows by setting active=true in manifest.csv",
            "Keep thresholds.csv tuned per target before strict gating",
            "Outcome: higher confidence without noisy false failures",
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main() -> None:
    out = Path("UnifiedFramework/DATA3/docs/validation/nightly/slides/nightly_ops_10_slide_deck.pptx")
    build_deck(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
