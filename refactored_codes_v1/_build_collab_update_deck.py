#!/usr/bin/env python3
"""Build the DATA3 collaborator-update deck (+ a separate backup deck).

Design choices (per request):
  * White background, dark (navy) text throughout — no dark slides.
  * Math/equations/units in EQUATION MODE: native, editable PowerPoint OMML
    (via _eq2omml.latex_to_omath -> pandoc), with a clean unicode fallback for
    non-PowerPoint renderers. Inline units in bullets use proper sub/superscripts.
  * Lean: 8 main slides ("with explanatory"); the 3 explanatory slides
    (Feedback, Topics, σ–B bridge) can be dropped for a 5-slide core
    (Title + Péclet + Donnan master + Taylor LaCl3 + Taylor CaCl2).
  * All other figure-pitch slides go to a SEPARATE backup file.

Run with pandoc on PATH:
  PATH="$HOME/miniforge3/bin:$PATH" python _build_collab_update_deck.py
"""
import json
import sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from xml.sax.saxutils import escape as _xesc

HERE = Path(__file__).resolve().parent
BUILD = HERE / "_collab_update_build"
sys.path.insert(0, str(HERE))
from _eq2omml import latex_to_omath  # noqa: E402

OUT_MAIN = HERE / "DATA3_collaborator_update_2026-06-24.pptx"
OUT_BACKUP = HERE / "DATA3_collaborator_update_2026-06-24_backup.pptx"

NAVY = RGBColor(0x0C, 0x23, 0x40)
GOLD = RGBColor(0xB8, 0x8A, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x20, 0x20, 0x20)
GRAY = RGBColor(0x66, 0x66, 0x66)
# Transport-mechanism semantic colors (carried from the source figure) + schematic tints
GREEN = RGBColor(0x2F, 0xA8, 0x4F)   # convection
BLUE = RGBColor(0x2E, 0x6F, 0xDB)    # diffusion
RED = RGBColor(0xCC, 0x3B, 0x33)     # electromigration / + ion
IONB = RGBColor(0x3B, 0x7D, 0xD8)    # − ion
LGRAY = RGBColor(0xF1, 0xF1, 0xF1)   # equation band
WALL = RGBColor(0x9C, 0x9C, 0x9C)    # membrane wall
PORE = RGBColor(0xEC, 0xF2, 0xFB)    # pore interior
GLOW = RGBColor(0xFC, 0xE2, 0xC4)    # low-molarity highlight
GOLDT = RGBColor(0xF8, 0xEF, 0xD6)   # callout tint
TITLE_FONT = "Cambria"
BODY_FONT = "Calibri"
SW, SH = Inches(13.333), Inches(7.5)
FOOTER_L = "DATA3 single-salt NF270 diafiltration  ·  collaborator update"
DATE = "2026-06-24"
caps = json.load(open(BUILD / "captions.json"))

# Native equations: latex + clean-unicode fallback (shown by LibreOffice/QA).
EQ = {
 "peclet":  (r"\mathrm{Pe}=\dfrac{\text{convection}}{\text{diffusion}}=\dfrac{J_w\,L}{D_m}",
             "Pe = convection / diffusion = Jw·L / Dₘ"),
 "reject":  (r"R=\dfrac{\sigma\left(1-e^{-\mathrm{Pe}}\right)}{1-\sigma\,e^{-\mathrm{Pe}}}",
             "R = σ(1 − e⁻ᴾᵉ)/(1 − σ e⁻ᴾᵉ)"),
 "donnan":  (r"B(c)=\dfrac{J_w\,P_0\left(-X+\sqrt{X^{2}+4k^{2}c^{2}}\right)}{2c}",
             "B(c) = J_w P₀(−X + √(X² + 4k²c²)) / (2c)"),
 "plateau": (r"B_{\text{plateau}}=J_w P_0 k,\quad c_{\text{knee}}=\dfrac{X}{2k},\quad u=\dfrac{c}{c_{\text{knee}}}=\dfrac{2kc}{X}",
             "B_plateau = J_w P₀ k,   c_knee = X/2k,   u = c/c_knee = 2kc/X"),
 "taylor":  (r"B(c)=\beta_0+\beta_1 c+\beta_2 c^{2}+\beta_3 c^{3}",
             "B(c) = β₀ + β₁c + β₂c² + β₃c³"),
 "satexp":  (r"B(c)=B_{\infty}\left(1-e^{-c/c^{*}}\right)",
             "B(c) = B_∞(1 − e⁻ᶜᐟᶜ*)"),
 "nernstplanck": (r"J_i = v_s\,C_i\,K_{i,v}\;-\;D_i\,\nabla C_i\;-\;z_i\,C_i\,D_i\,\dfrac{F}{RT}\,\nabla\psi",
                  "J_i = v_s·C_i·K_i,v  −  D_i·∇C_i  −  z_i·C_i·D_i·(F/RT)·∇ψ"),
}

# Footer references per equation slide — matched to Architecture.md and the DATA1/DATA2 papers.
#   DATA1 = Ouimet et al., J. Membr. Sci. 641 (2022) 119743  (solution-diffusion framework, constant B)
#   DATA2 = Liu et al., Ind. Eng. Chem. Res. 64 (2025) 12111 (derives Pe = JwL/Dm, eq 25; polynomial B(c), eq 31)
REF = {
 "peclet":  "Refs — Pe = JwL/Dm: DATA2, Liu et al., Ind. Eng. Chem. Res. 64 (2025) 12111 (eq 25); Spiegler–Kedem rejection: Spiegler & Kedem, Desalination 1 (1966) 311; solution-diffusion (Lp, B, σ): Kedem & Katchalsky, Biochim. Biophys. Acta 27 (1958) 229; framework: Ouimet et al., J. Membr. Sci. 641 (2022) 119743 (DATA1).",
 "donnan":  "Refs — Donnan–dielectric B(c): Yaroshchuk, Adv. Colloid Interface Sci. 85 (2000) 193; Bandini & Vezzani, Chem. Eng. Sci. 58 (2003) 3303 (DSPM-DE); Donnan-equilibrium partition cf. Bolt, J. Colloid Sci. 10 (1955) 206 (as in DATA2).",
 "taylor":  "Refs — empirical B(c): polynomial (multiplicative) form per DATA2, Liu et al., Ind. Eng. Chem. Res. 64 (2025) 12111 (eq 31); saturating-exp & Donnan-bounded forms: this work (DATA3).",
}

_SUB = {**{c: d for c, d in zip("₀₁₂₃₄₅₆₇₈₉", "0123456789")},
        "ᵢ": "i", "ₙ": "n", "ⱼ": "j", "ᵥ": "v", "ᵤ": "u", "ᵣ": "r", "ₐ": "a",
        "ₑ": "e", "ₒ": "o", "ₓ": "x", "ₚ": "p", "ₛ": "s", "ₜ": "t", "ₕ": "h",
        "ₖ": "k", "ₗ": "l", "ₘ": "m"}
_SUP = {**{c: d for c, d in zip("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")},
        "⁻": "-", "⁺": "+", "ⁱ": "i", "ⁿ": "n", "²": "2", "³": "3", "¹": "1"}


def _solid(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()


def bg(slide, color=WHITE):
    r = slide.shapes.add_shape(1, 0, 0, SW, SH); _solid(r, color)
    slide.shapes._spTree.remove(r._element); slide.shapes._spTree.insert(2, r._element)


def tbx(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    # Accept either EMU (Inches(...), an int subclass) or raw inches as floats.
    _e = lambda v: v if isinstance(v, int) else Inches(v)
    b = slide.shapes.add_textbox(_e(l), _e(t), _e(w), _e(h)); tf = b.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    return b, tf


def run(p, text, size, color, bold=False, italic=False, font=BODY_FONT):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    r.font.italic = italic; r.font.name = font
    return r


def rich(p, text, size, color, bold=False, font=BODY_FONT):
    buf, mode = "", 0

    def flush(m):
        nonlocal buf
        if not buf:
            return
        r = run(p, buf, size, color, bold=bold, font=font)
        if m == -1:
            r.font._rPr.set("baseline", "-25000")
        elif m == 1:
            r.font._rPr.set("baseline", "30000")
        buf = ""
    for ch in text:
        if ch in _SUB:
            m, o = -1, _SUB[ch]
        elif ch in _SUP:
            m, o = 1, _SUP[ch]
        else:
            m, o = 0, ch
        if m != mode:
            flush(mode); mode = m
        buf += o
    flush(mode)


def footer(slide, n, refs=None):
    if refs:   # small references line(s) just above the standard footer
        b0, tf0 = tbx(slide, Inches(0.4), Inches(6.74), Inches(12.55), Inches(0.30))
        run(tf0.paragraphs[0], refs, 8, GRAY, italic=True)
    b, tf = tbx(slide, Inches(0.4), Inches(7.07), Inches(9.5), Inches(0.32))
    run(tf.paragraphs[0], FOOTER_L, 9, GRAY)
    b2, tf2 = tbx(slide, Inches(10.0), Inches(7.07), Inches(2.93), Inches(0.32))
    tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run(tf2.paragraphs[0], f"{DATE}  ·  {n}", 9, GRAY)


def title(slide, text):
    b, tf = tbx(slide, Inches(0.5), Inches(0.28), Inches(12.33), Inches(0.85), MSO_ANCHOR.MIDDLE)
    rich(tf.paragraphs[0], text, 26, NAVY, bold=True, font=TITLE_FONT)


def place_fit(slide, path, L, T, W, H):
    """Place image scaled to fit (W,H) preserving aspect ratio, centered.
    Returns (left, top, w, h) in EMU so callers can flow content below it."""
    iw, ih = Image.open(path).size
    s = min(W / iw, H / ih); w, h = int(iw * s), int(ih * s)
    left, top = L + (W - w) // 2, T + (H - h) // 2
    slide.shapes.add_picture(str(path), left, top, w, h)
    return left, top, w, h


def place_eq(slide, key, y, sz=18):
    """Native, editable OMML equation centered horizontally; unicode fallback."""
    latex, fallback = EQ[key]
    b, tf = tbx(slide, Inches(0.6), Inches(y), Inches(12.13), Inches(0.55))
    p = tf.paragraphs[0]            # do NOT also set p.alignment — we inject one <a:pPr> ourselves
    sz100 = int(sz * 100)
    try:
        omath = latex_to_omath(latex)
        ppr = ('<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" algn="ctr">'
               f'<a:defRPr sz="{sz100}"><a:solidFill><a:srgbClr val="0C2340"/></a:solidFill></a:defRPr></a:pPr>')
        ac = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
              '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">'
              '<a14:m xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
              f'<m:oMathPara>{omath}</m:oMathPara></a14:m></mc:Choice>'
              '<mc:Fallback xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
              f'<a:r><a:rPr lang="en-US" sz="{sz100}" dirty="0"><a:solidFill><a:srgbClr val="0C2340"/></a:solidFill>'
              f'<a:latin typeface="Cambria Math"/></a:rPr><a:t>{_xesc(fallback)}</a:t></a:r></mc:Fallback>'
              '</mc:AlternateContent>')
        p._p.insert(0, parse_xml(ppr)); p._p.append(parse_xml(ac))
    except Exception as exc:
        print(f"  [eq fallback] {key}: {exc!r}")
        p.alignment = PP_ALIGN.CENTER
        rich(p, fallback, sz, NAVY, font="Cambria")


def bullets(slide, items, L, T, W, H, size=14, gap=6):
    b, tf = tbx(slide, L, T, W, H, MSO_ANCHOR.TOP)
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr(); pPr.set("indent", "-228600"); pPr.set("marL", "228600")
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        if lead:
            rich(p, lead, size, NAVY, bold=True)
        if rest:
            rich(p, (" " if lead else "") + rest, size, DARK)


def ashape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0):
    """Native autoshape (rectangle / rounded-rect / oval / arrow), flat (no shadow),
    fully editable in PowerPoint. Coords in inches."""
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    return sp


def shape_text(sp, text, size, color, bold=True, align=None, font=BODY_FONT):
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align or PP_ALIGN.CENTER
    run(p, text, size, color, bold=bold, font=font)


def _ions(slide, x, w, n, row_ys, dia=0.16):
    """Lay n ions (alternating + red / − blue) across two rows inside a pore."""
    cols = max(1, (n + 1) // 2)
    xs = [x + (i + 0.5) * w / cols - dia / 2 for i in range(cols)]
    k = 0
    for r, ry in enumerate(row_ys):
        for i, xx in enumerate(xs):
            if k >= n:
                break
            plus = ((i + r) % 2 == 0)
            ov = ashape(slide, MSO_SHAPE.OVAL, xx, ry, dia, dia, fill=(RED if plus else IONB))
            shape_text(ov, "+" if plus else "–", 8, WHITE, bold=True)
            k += 1


def nernst_planck_slide(bk, n):
    """Backup slide: 'Shifting Transport Dominance — Extended Nernst–Planck',
    rebuilt entirely from native, editable PowerPoint objects (equation, bars,
    pore schematics, ions, arrows, callout) in our white/navy/gold format."""
    s = bk.slides.add_slide(bk.slide_layouts[6]); bg(s)
    title(s, "Backup — Shifting Transport Dominance (Extended Nernst–Planck)")

    # --- equation band (native OMML equation + colored term underlines) ---
    ashape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 1.02, 12.33, 1.04, fill=LGRAY)
    place_eq(s, "nernstplanck", 1.10, sz=17)
    for lab, col, cx in (("Convection", GREEN, 3.15), ("Diffusion", BLUE, 6.55),
                         ("Electromigration", RED, 9.95)):
        ashape(s, MSO_SHAPE.RECTANGLE, cx - 0.75, 1.72, 1.5, 0.035, fill=col)
        b, tf = tbx(s, cx - 1.15, 1.76, 2.3, 0.26); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run(tf.paragraphs[0], lab, 11, col, bold=True)

    cols = [
        dict(left=0.60, hdr="Low Molarity  (~0.01–0.1 M)", pct=(10, 20, 70), nions=6, glow=True,
             leg=[("Electromigration", "dominant", RED), ("Diffusion", "secondary", BLUE),
                  ("Convection", "minor", GREEN)]),
        dict(left=4.73, hdr="Moderate Molarity  (~0.1–0.5 M)", pct=(20, 35, 45), nions=10, glow=False,
             leg=[("Convection", "increasing", GREEN), ("Diffusion", "comparable", BLUE),
                  ("Electromigration", "significant", RED)]),
        dict(left=8.86, hdr="High Molarity  (~0.5–2+ M)", pct=(60, 30, 10), nions=16, glow=False,
             leg=[("Convection", "dominant", GREEN), ("Diffusion", "secondary", BLUE),
                  ("Electromigration", "minor", RED)]),
    ]
    CW = 3.87
    for c in cols:
        L = c["left"]
        b, tf = tbx(s, L, 2.20, CW, 0.30); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run(tf.paragraphs[0], c["hdr"], 12.5, NAVY, bold=True, font=TITLE_FONT)
        b, tf = tbx(s, L, 2.48, CW, 0.24); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        rich(tf.paragraphs[0], "Contributions to flux (Jᵢ)", 9.5, GRAY)

        # 100% stacked bar (native rectangles, each labeled)
        bx, bw, by, bh = L + 0.3, CW - 0.6, 2.74, 0.32
        x = bx
        for pct, scol in zip(c["pct"], (GREEN, BLUE, RED)):
            w = bw * pct / 100.0
            seg = ashape(s, MSO_SHAPE.RECTANGLE, x, by, w, bh, fill=scol, line=WHITE, lw=0.75)
            shape_text(seg, f"{pct}%", (10 if w > 1.0 else 9 if w > 0.55 else 7), WHITE, bold=True)
            x += w
        b, tf = tbx(s, L, by + bh + 0.01, CW, 0.18); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run(tf.paragraphs[0], "Contribution (%)", 8.5, GRAY)

        # pore schematic
        sy, sh = 3.30, 1.46
        px, pw = L + 0.32, CW - 0.64
        if c["glow"]:
            ashape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px - 0.10, sy - 0.10, pw + 0.20, sh + 0.20, fill=GLOW)
        ashape(s, MSO_SHAPE.RECTANGLE, px, sy + 0.16, pw, sh - 0.32, fill=PORE, line=WALL, lw=0.75)
        ashape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px, sy, pw, 0.16, fill=WALL)
        ashape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px, sy + sh - 0.16, pw, 0.16, fill=WALL)
        for wy in (sy + 0.005, sy + sh - 0.155):
            b, tf = tbx(s, px, wy, pw, 0.16); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            run(tf.paragraphs[0], "–     –     –     –     –     –", 9, WHITE, bold=True)
        _ions(s, px + 0.08, pw - 0.16, c["nions"], [sy + 0.19, sy + 1.14])
        for pct, acol, yy in zip(c["pct"], (GREEN, BLUE, RED), (3.74, 3.94, 4.14)):
            alen = max(0.30, (pw - 0.55) * pct / 100.0)
            ashape(s, MSO_SHAPE.RIGHT_ARROW, px + 0.22, yy, alen, 0.13, fill=acol)
        ashape(s, MSO_SHAPE.RIGHT_ARROW, L + 0.02, sy + sh / 2 - 0.065, 0.22, 0.13, fill=DARK)
        ashape(s, MSO_SHAPE.RIGHT_ARROW, px + pw + 0.02, sy + sh / 2 - 0.065, 0.22, 0.13, fill=DARK)

        # per-column legend
        ly = 5.04
        for name, desc, lcol in c["leg"]:
            ashape(s, MSO_SHAPE.RIGHT_ARROW, L + 0.15, ly + 0.03, 0.26, 0.13, fill=lcol)
            b, tf = tbx(s, L + 0.5, ly - 0.02, CW - 0.55, 0.24)
            p = tf.paragraphs[0]; run(p, name, 10, NAVY, bold=True); run(p, f"  ({desc})", 10, DARK)
            ly += 0.235

    b, tf = tbx(s, 0.5, 4.80, 12.33, 0.20); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    rich(tf.paragraphs[0], "Pore cross-section — fixed (–) wall charges, mobile ions (+/–), flux arrows sized by contribution;  flow is Feed → Permeate (left → right)", 9, GRAY)

    # --- Donnan-potential callout ---
    ashape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 5.90, 12.33, 1.04, fill=GOLDT, line=GOLD, lw=1.0)
    ic = ashape(s, MSO_SHAPE.OVAL, 0.74, 6.10, 0.62, 0.62, fill=WHITE, line=GOLD, lw=1.25)
    shape_text(ic, "ψ", 19, NAVY, bold=False, font=TITLE_FONT)
    b, tf = tbx(s, 1.55, 5.97, 11.1, 0.92)
    p = tf.paragraphs[0]
    run(p, "Role of fixed negative charges & the Donnan potential.  ", 12, NAVY, bold=True, font=TITLE_FONT)
    run(p, "Negatively charged pore walls repel co-ions and attract counter-ions, creating a Donnan "
           "potential ψD that drives electromigration. At low molarity the low ionic strength magnifies "
           "ψD, making electromigration the primary driver; as molarity rises, charge screening by "
           "abundant ions shrinks ψD, so convection increasingly governs transport.", 11, DARK)
    footer(s, n)
    return s


def figure_slide(prs, n, ttl, img, eq_keys, items, where=None, refs=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title(s, ttl)
    # All these figures are wide (aspect 2–3); let them fill the width. Cap the
    # height so the equations + bullets still fit, more room when fewer equations.
    h_cap = {0: 4.55, 1: 4.10}.get(len(eq_keys), 3.70)
    if where:                                  # reserve room for the variable glossary + refs line
        h_cap -= 0.85
    _, top, _, h = place_fit(s, BUILD / img, Inches(0.32), Inches(1.05), Inches(12.70), Inches(h_cap))
    y = (top + h) / 914400.0 + 0.12            # flow content from the figure's true bottom
    for k in eq_keys:
        place_eq(s, k, y); y += 0.52
    if where:                                  # "where ..." definitions of every variable above
        b, tf = tbx(s, Inches(0.7), Inches(y + 0.03), Inches(11.93), Inches(0.64))
        p = tf.paragraphs[0]
        run(p, "where  ", 11, NAVY, bold=True, italic=True)
        rich(p, where, 11, DARK)
        y += 0.70
    bull_bottom = 6.58 if refs else 7.0
    bullets(s, items, Inches(0.6), Inches(y + 0.06), Inches(12.13),
            Inches(max(0.5, bull_bottom - (y + 0.06))), size=(13 if where else 14), gap=(4 if where else 5))
    footer(s, n, refs=refs)
    return s


# ============================================================== MAIN DECK
prs = Presentation(); prs.slide_width = SW; prs.slide_height = SH
B6 = prs.slide_layouts[6]

# 1 — Title (white bg, dark text)
s = prs.slides.add_slide(B6); bg(s)
b, tf = tbx(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2))
rich(tf.paragraphs[0], "DATA3: Single-Salt NF270 Diafiltration", 40, NAVY, bold=True, font=TITLE_FONT)
p = tf.add_paragraph(); p.space_before = Pt(10)
rich(p, "Concentration-dependent solute permeability & the Donnan–dielectric model", 20, GOLD)
b, tf = tbx(s, Inches(0.95), Inches(5.2), Inches(11), Inches(1.2))
rich(tf.paragraphs[0], "Collaborator update  ·  June 24, 2026", 16, NAVY, bold=True)
p = tf.add_paragraph(); p.space_before = Pt(4)
rich(p, "Keshav Kasturi Rangan  ·  University of Notre Dame", 14, GRAY)
logo = HERE / "notre_dame_logo.png"
if logo.exists():
    place_fit(s, logo, Inches(10.9), Inches(0.45), Inches(2.0), Inches(1.25))

# 2 — Feedback from Friday (explanatory)
s = prs.slides.add_slide(B6); bg(s); title(s, "Feedback from Friday — addressed")
bullets(s, [
    ("Time-series correction:", "permeate ICP now anchored at vial close (DATA2 convention); mass, retentate, and permeate share one time axis — the invented tube-transit placement is retired."),
    ("Identifiability contours:", "σ×B log-WSSE surfaces generated per response (mass / permeate / retentate) at the identified Lₚ."),
    ("Color coding:", "fit/data plots follow the DATA2 convention (retentate vs permeate vs model)."),
    ("Scope:", "all changes are DATA3 / NF270-only; DATA1, DATA2, and the legacy MATLAB are untouched."),
], Inches(0.7), Inches(1.5), Inches(12), Inches(5.0), size=18, gap=12)
footer(s, 2)

# 3 — Topics I need help with (explanatory)
s = prs.slides.add_slide(B6); bg(s); title(s, "Topics I need help with")
bullets(s, [
    ("Prioritization:", "which question below is the right one to chase first?"),
    ("Donnan–dielectric equation:", "the partition form, its assumptions, and the derivation of B(c)."),
    ("Interfacial concentration:", "what range of membrane-interface concentration do the experiments actually probe?"),
    ("Péclet framing:", "the Péclet number and the rejection law R(σ, Pe) — is this the right lens?"),
    ("Transport components:", "dielectric exclusion, cross-coefficients Dᵢⱼ, B(c), and activity coefficients — which to include?"),
    ("Identifiability:", "given σ is unidentified single-experiment, how far can B(c) claims go?"),
], Inches(0.7), Inches(1.45), Inches(12), Inches(5.2), size=17, gap=10)
footer(s, 3)

# 4 — σ–B bridge (explanatory)
figure_slide(prs, 4, "Single experiments determine B, not σ", "figs/all_slide10_img0.png", [],
    [("Each response’s log-WSSE has a narrow valley in B but is nearly flat in σ", "(representative concentrating CaCl₂; Lₚ fixed at its identified optimum)."),
     ("σ rails to its bound while B stays well-determined", "— the two are correlated, so one experiment cannot recover σ."),
     ("Consequence:", "“is B concentration-dependent?” must be posed carefully — we lean on mechanism, not fit alone.")])

# 5 — Péclet regimes (core)
figure_slide(prs, 5, "The same membrane spans diffusion- and convection-limited regimes",
    "figs/slide13_img0.png", ["peclet", "reject"],
    [("Diluting NaCl is diffusion-limited (Pe* ≈ 10⁻³); concentrating NaCl / CaCl₂ / LaCl₃ are convection-dominated (Pe* ≈ 77 / 36 / 36).", ""),
     ("Mechanistic inference, not a direct demonstration:", "Pe and B come from the same fitted model — this motivates B(c), it does not prove it."),
     ("To tabulate for review:", "water flux Jw, membrane length-scale L, and diffusivity Dₘ per salt.")],
    where="Pe — Péclet number;  Jw — water (volume) flux through the membrane;  L — membrane thickness;  "
          "Dₘ — solute diffusivity in the membrane;  R — salt rejection;  "
          "σ — reflection coefficient (0 = no rejection, 1 = complete).",
    refs=REF["peclet"])

# 6 — Donnan master curve (core)
figure_slide(prs, 6, "All single-salt windows sit on the Donnan plateau",
    "figs/slide14_img0.png", ["donnan", "plateau"],
    [("B(c) collapses onto one master curve; every measured window is at u ≈ 10³–10⁵ (screened plateau)", "→ B is ≈ constant in-window (the screened-plateau value)."),
     ("The knee c_knee = X / 2k rails 10³–10⁵× below the data", "→ the fixed-charge scale X is not identified from one salt.")],
    where="B(c) — solute permeability at interfacial concentration c;  c — membrane-interface salt concentration;  "
          "Jw — water flux;  P₀ — intrinsic permeability prefactor;  X — effective membrane fixed-charge density;  "
          "k — lumped Donnan–dielectric partition factor;  B_plateau — high-c plateau value;  "
          "c_knee = X/2k — regime knee;  u = c/c_knee — screening parameter.",
    refs=REF["donnan"])

# 7 — Taylor vs Donnan, LaCl3 (core)
figure_slide(prs, 7, "B(c): in-window agreement, out-of-window divergence — LaCl₃",
    "figs/slide15_crop.png", ["taylor", "satexp"],
    [("Within the measured window, constant / linear / quadratic / cubic / saturating-exp all agree (B ≈ 0.3 µm·s⁻¹).", ""),
     ("Extrapolated, the polynomials diverge unphysically (cubic → negative); Donnan and saturating-exp stay bounded.", ""),
     ("The Donnan–dielectric is the mechanistically-grounded, physically-bounded form", "— the data do not discriminate it from a constant in-window.")],
    where="B(c) — solute permeability;  c — interfacial concentration;  β₀–β₃ — polynomial (Taylor) expansion "
          "coefficients;  B∞ — saturating permeability as c → ∞;  c* — characteristic saturation concentration.",
    refs=REF["taylor"])

# 8 — Taylor vs Donnan, CaCl2 (core)
figure_slide(prs, 8, "Same behavior for CaCl₂", "figs/slide16_crop.png", [],
    [("Identical picture: in-window agreement, out-of-window polynomial divergence, Donnan bounded.", ""),
     ("On parsimony (AIC) a constant fits as well in-window", "— prefer Donnan for mechanism and physical extrapolation, not in-window fit."),
     ("Next step to pin X:", "a joint multi-salt fit (re-parameterized by ionic strength) reaches lower u and probes the knee.")])

prs.save(str(OUT_MAIN))
print(f"SAVED MAIN {OUT_MAIN}  ({len(prs.slides._sldIdLst)} slides)")


# ============================================================== BACKUP DECK
bk = Presentation(); bk.slide_width = SW; bk.slide_height = SH
s = bk.slides.add_slide(bk.slide_layouts[6]); bg(s)
b, tf = tbx(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.6), MSO_ANCHOR.MIDDLE)
rich(tf.paragraphs[0], "Backup — Supporting Figures", 36, NAVY, bold=True, font=TITLE_FONT)
p = tf.add_paragraph(); p.space_before = Pt(6)
rich(p, "Fits, full σ×B contours, Péclet regression, NaCl B(c), concentration–flux", 15, GRAY)


def short_title(cap):
    cap = " ".join(cap.replace("\n", " ").split())
    parts = cap.split(".")
    body = (parts[0].strip() + ". " + parts[1].strip()) if len(parts) >= 2 and parts[0].strip().lower().startswith("figure") else cap
    if len(body) > 80:
        body = body[:80].rsplit(" ", 1)[0].rstrip(" ,;:-") + "…"
    return body


BACKUP = [2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 17, 18]
for i, sl in enumerate(BACKUP, start=2):
    s = bk.slides.add_slide(bk.slide_layouts[6]); bg(s)
    title(s, "Backup — " + short_title(caps[str(sl)]))
    imgs = sorted(BUILD.glob(f"figs/all_slide{sl}_img*.png"))
    if len(imgs) == 1:
        place_fit(s, imgs[0], Inches(0.5), Inches(1.2), Inches(12.33), Inches(5.6))
    elif len(imgs) >= 2:
        asp = [Image.open(p).size[0] / Image.open(p).size[1] for p in imgs[:2]]
        if all(a > 2.5 for a in asp):   # both very wide -> stack full width (don't squish side-by-side)
            place_fit(s, imgs[0], Inches(0.6), Inches(1.3), Inches(12.13), Inches(2.75))
            place_fit(s, imgs[1], Inches(0.6), Inches(4.15), Inches(12.13), Inches(2.75))
        else:                            # mixed/square -> side by side, vertically centered
            place_fit(s, imgs[0], Inches(0.4), Inches(1.35), Inches(6.35), Inches(5.4))
            place_fit(s, imgs[1], Inches(6.8), Inches(1.35), Inches(6.35), Inches(5.4))
    footer(s, i)

# Native (fully editable) recreation of the Extended Nernst–Planck transport figure
nernst_planck_slide(bk, len(bk.slides._sldIdLst) + 1)

bk.save(str(OUT_BACKUP))
print(f"SAVED BACKUP {OUT_BACKUP}  ({len(bk.slides._sldIdLst)} slides)")
